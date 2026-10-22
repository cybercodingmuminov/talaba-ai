import asyncio
import hashlib
import hmac
import json
import logging
import os
import sqlite3
import time
import uuid
from contextlib import closing
from datetime import date
from pathlib import Path
from urllib.parse import parse_qsl

import httpx
from docx import Document
from dotenv import load_dotenv
from fastapi import FastAPI, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Pt
from pydantic import BaseModel

load_dotenv()
BOT_TOKEN = os.environ["BOT_TOKEN"]
PUBLIC_URL = os.environ.get("PUBLIC_URL", "http://localhost:8000").rstrip("/")
ORIGINS = [o.strip() for o in os.environ.get("ALLOWED_ORIGINS", "*").split(",")]
DAILY_LIMIT = int(os.environ.get("DAILY_LIMIT", "3"))  # bir foydalanuvchiga kuniga
GLOBAL_LIMIT = int(os.environ.get("GLOBAL_LIMIT", "100"))  # hamma uchun kuniga (bepul kvota tugamasligi uchun)
GEMINI_API_KEY = os.environ["GEMINI_API_KEY"]
MODEL = os.environ.get("GEMINI_MODEL", "gemini-flash-latest")
DEV = os.environ.get("DEV", "0") == "1"  # True bo'lsa Telegramsiz sinash mumkin

log = logging.getLogger("talaba")
http = httpx.AsyncClient(timeout=120)
gate = asyncio.Semaphore(3)  # bir vaqtda ko'pi bilan 3 ta AI so'rovi

FILES = Path("files")
FILES.mkdir(exist_ok=True)
DB = "usage.db"

app = FastAPI()
app.add_middleware(CORSMiddleware, allow_origins=ORIGINS, allow_methods=["POST"], allow_headers=["*"])
app.mount("/files", StaticFiles(directory=FILES), name="files")


# ---------- Telegram autentifikatsiya ----------
def verify(init_data: str) -> int:
    if not init_data:
        if DEV:
            return 0
        raise HTTPException(401, "Ilovani Telegram ichida oching.")
    data = dict(parse_qsl(init_data, keep_blank_values=True))
    got = data.pop("hash", "")
    check = "\n".join(f"{k}={v}" for k, v in sorted(data.items()))
    secret = hmac.new(b"WebAppData", BOT_TOKEN.encode(), hashlib.sha256).digest()
    calc = hmac.new(secret, check.encode(), hashlib.sha256).hexdigest()
    if not hmac.compare_digest(calc, got):
        raise HTTPException(401, "Autentifikatsiya xatosi.")
    if time.time() - int(data.get("auth_date", "0")) > 86400:
        raise HTTPException(401, "Sessiya eskirgan, ilovani qayta oching.")
    return int(json.loads(data["user"])["id"])


# ---------- Kunlik limit ----------
def db():
    c = sqlite3.connect(DB)
    c.execute("CREATE TABLE IF NOT EXISTS usage(uid INTEGER, day TEXT, n INTEGER, PRIMARY KEY(uid, day))")
    return c


def take_slot(uid: int, limit: int) -> bool:
    day = date.today().isoformat()
    with closing(db()) as c:
        c.execute("INSERT OR IGNORE INTO usage VALUES(?,?,0)", (uid, day))
        cur = c.execute("UPDATE usage SET n=n+1 WHERE uid=? AND day=? AND n<?", (uid, day, limit))
        c.commit()
        return cur.rowcount == 1


def give_back(uid: int):
    with closing(db()) as c:
        c.execute("UPDATE usage SET n=n-1 WHERE uid=? AND day=? AND n>0", (uid, date.today().isoformat()))
        c.commit()


def cleanup():
    now = time.time()
    for f in FILES.iterdir():
        if now - f.stat().st_mtime > 86400:
            f.unlink(missing_ok=True)


# ---------- AI so'rovi ----------
def level(size: str) -> int:
    return 0 if size.startswith("Qisqa") else 2 if size.startswith("Keng") else 1


def make_prompt(service: str, topic: str, lang: str, size: str):
    L = level(size)
    system = (
        "Sen talabalar uchun o'quv materiallari tayyorlaydigan yordamchisan. "
        f"Javob tili: {lang}. Faqat bitta JSON obyekt qaytar — izoh, markdown yoki ``` belgilarisiz. "
        "Matematik formulalarni oddiy matn ko'rinishida yoz (LaTeX ishlatma)."
    )
    t = f'Mavzu: "{topic}". '
    if service == "slides":
        n = [8, 13, 20][L]
        u = t + f"{n} ta slayddan iborat prezentatsiya tuz. Format: " + '{"title": str, "slides": [{"title": str, "bullets": [str]}]}' + \
            ". Har slaydda 3-5 ta qisqa band (har biri 15 so'zdan oshmasin). Oxirgi slayd — xulosa."
    elif service in ("coursework", "project"):
        n, w = [4, 6, 9][L], [200, 300, 400][L]
        what = "referat / mustaqil ish" if service == "coursework" else "individual loyiha rejasi (maqsad, bosqichlar, texnologiyalar, kutilgan natija)"
        u = t + f"{what} tayyorla. {n} ta bo'lim, har biri kamida {w} so'z. Format: " + \
            '{"title": str, "sections": [{"heading": str, "text": str}], "references": [str]}' + \
            ". Bo'lim matnida abzaslarni \\n\\n bilan ajrat. Adabiyotlar ro'yxatida faqat mavjudligiga ishonchli manbalarni keltir."
    elif service in ("test", "math"):
        n = [10, 20, 30][L]
        kind = "matematik masalalardan iborat test" if service == "math" else "test"
        u = t + f"{n} ta savoldan iborat {kind} tuz. Format: " + \
            '{"title": str, "questions": [{"q": str, "options": [str, str, str, str], "answer": int}]}' + \
            ". answer — to'g'ri variantning indeksi (0-3). Variantlar mazmunan farqli bo'lsin."
    elif service == "cross":
        n = [10, 15, 20][L]
        u = t + f"{n} ta so'zdan krossvord uchun so'zlar va ta'riflar tayyorla. Format: " + \
            '{"title": str, "entries": [{"word": str, "clue": str}]}' + ". So'zlar bitta so'zdan iborat, bosh harflarsiz."
    elif service == "learn":
        u = t + 'Qisqa dars tayyorla: kirish, 5 ta asosiy fikr, oxirida javoblari bilan 3 ta tekshiruv savoli. Format: {"text": str}. Oddiy matn, abzaslar \\n\\n bilan.'
    else:
        raise HTTPException(400, "Noma'lum xizmat.")
    return system, u


async def ask(system: str, user: str) -> dict:
    body = {
        "systemInstruction": {"parts": [{"text": system}]},
        "contents": [{"role": "user", "parts": [{"text": user}]}],
        "generationConfig": {"responseMimeType": "application/json", "maxOutputTokens": 8192},
    }
    url = f"https://generativelanguage.googleapis.com/v1beta/models/{MODEL}:generateContent"
    async with gate:
        r = await http.post(url, json=body, headers={"x-goog-api-key": GEMINI_API_KEY})
    if r.status_code == 429:
        raise HTTPException(503, "AI hozir band yoki bepul kvota tugadi. Birozdan keyin urinib ko'ring.")
    r.raise_for_status()
    parts = r.json()["candidates"][0]["content"]["parts"]
    text = "".join(p.get("text", "") for p in parts)
    return json.loads(text[text.index("{"): text.rindex("}") + 1])


# ---------- Fayl yaratish ----------
def build_pptx(d: dict, path: Path):
    ink, acc, white = RGBColor(0x14, 0x21, 0x3D), RGBColor(0x0E, 0x9F, 0x8E), RGBColor(255, 255, 255)
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = ink
    s.shapes.title.text = d["title"]
    s.placeholders[1].text = "AI Shaxsiy Talaba Platformasi"
    for ph in (s.shapes.title, s.placeholders[1]):
        for p in ph.text_frame.paragraphs:
            p.font.color.rgb = white
    s.shapes.title.text_frame.paragraphs[0].font.size = Pt(38)
    for sl in d["slides"]:
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = sl["title"]
        tp = s.shapes.title.text_frame.paragraphs[0]
        tp.font.size, tp.font.bold, tp.font.color.rgb = Pt(32), True, acc
        tf = s.placeholders[1].text_frame
        tf.word_wrap = True
        for i, b in enumerate(sl["bullets"]):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = b
            p.font.size = Pt(22)
            p.font.color.rgb = ink
    prs.save(path)


def build_docx(d: dict, path: Path, kind: str):
    doc = Document()
    doc.add_heading(d["title"], 0)
    if kind == "work":
        doc.add_heading("Reja", 1)
        for i, s in enumerate(d["sections"], 1):
            doc.add_paragraph(f"{i}. {s['heading']}")
        for s in d["sections"]:
            doc.add_heading(s["heading"], 1)
            for para in s["text"].split("\n\n"):
                if para.strip():
                    doc.add_paragraph(para.strip())
        if d.get("references"):
            doc.add_heading("Foydalanilgan adabiyotlar", 1)
            for i, r in enumerate(d["references"], 1):
                doc.add_paragraph(f"{i}. {r}")
    elif kind == "test":
        letters = "ABCD"
        for i, q in enumerate(d["questions"], 1):
            doc.add_paragraph(f"{i}. {q['q']}")
            for j, o in enumerate(q["options"][:4]):
                doc.add_paragraph(f"    {letters[j]}) {o}")
        doc.add_page_break()
        doc.add_heading("Javoblar", 1)
        doc.add_paragraph(", ".join(f"{i}-{letters[int(q['answer']) % 4]}" for i, q in enumerate(d["questions"], 1)))
    elif kind == "cross":
        doc.add_heading("Ta'riflar", 1)
        for i, e in enumerate(d["entries"], 1):
            doc.add_paragraph(f"{i}. {e['clue']}  ({len(e['word'])} harf)")
        doc.add_page_break()
        doc.add_heading("Javoblar", 1)
        for i, e in enumerate(d["entries"], 1):
            doc.add_paragraph(f"{i}. {e['word'].upper()}")
    doc.save(path)


KIND = {"coursework": "work", "project": "work", "test": "test", "math": "test", "cross": "cross"}


# ---------- Endpointlar ----------
class Req(BaseModel):
    service: str
    topic: str
    lang: str = "O'zbekcha"
    size: str = ""
    initData: str = ""


@app.get("/")
def health():
    return {"ok": True}


@app.post("/api/generate")
async def generate(r: Req):
    uid = verify(r.initData)
    topic = r.topic.strip()
    if not 3 <= len(topic) <= 300:
        raise HTTPException(400, "Mavzu 3 dan 300 belgigacha bo'lishi kerak.")
    system, user = make_prompt(r.service, topic, r.lang[:20], r.size)
    if not take_slot(-1, GLOBAL_LIMIT):
        raise HTTPException(503, "Bugungi umumiy limit tugadi. Ertaga qayta urinib ko'ring.")
    if not take_slot(uid, DAILY_LIMIT):
        give_back(-1)
        raise HTTPException(429, f"Kunlik limit ({DAILY_LIMIT} ta) tugadi. Ertaga qayta urinib ko'ring.")
    try:
        cleanup()
        data = await ask(system, user)
        if r.service == "learn":
            return {"text": data["text"]}
        ext = "pptx" if r.service == "slides" else "docx"
        path = FILES / f"{uuid.uuid4().hex}.{ext}"
        if r.service == "slides":
            await asyncio.to_thread(build_pptx, data, path)
        else:
            await asyncio.to_thread(build_docx, data, path, KIND[r.service])
        return {"url": f"{PUBLIC_URL}/files/{path.name}"}
    except HTTPException:
        give_back(uid)
        give_back(-1)
        raise
    except Exception:
        log.exception("generatsiya xatosi")
        give_back(uid)
        give_back(-1)
        raise HTTPException(500, "Xatolik yuz berdi. Qayta urinib ko'ring.")
